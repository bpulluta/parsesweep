#!/usr/bin/env python3
"""
Create a styled Excel documentation file from the extraction schema.
Shows field names and descriptions in an easy-to-read format.
"""

import json
import pandas as pd
from pathlib import Path
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

WORKSPACE_ROOT = Path("/Users/bpulluta/StreamlineExtract")
SCHEMA_FILE = WORKSPACE_ROOT / "schemas" / "air_quality_permits_schema.json"


def load_schema() -> dict:
    """Load the schema JSON file."""
    return json.loads(SCHEMA_FILE.read_text())


def extract_field_info(schema_properties: dict, parent_path: str = "") -> list:
    """Recursively extract field names and descriptions from schema."""
    rows = []
    
    for field_name, field_def in schema_properties.items():
        # Construct full path
        full_path = f"{parent_path}.{field_name}" if parent_path else field_name
        
        # Get description
        description = field_def.get("description", "")
        
        # Get type information
        field_type = ""
        if "type" in field_def:
            if isinstance(field_def["type"], list):
                field_type = " | ".join(field_def["type"])
            else:
                field_type = field_def["type"]
        elif "anyOf" in field_def:
            types = [t.get("type", "") for t in field_def["anyOf"] if "type" in t]
            field_type = " | ".join(filter(None, types))
        
        # Handle enum values
        enum_values = ""
        if "enum" in field_def:
            enum_values = ", ".join([f'"{v}"' for v in field_def["enum"]])
        
        # Check if required
        is_required = field_name in schema_properties.get("required", [])
        
        rows.append({
            "field_name": field_name,
            "full_path": full_path,
            "type": field_type,
            "enum": enum_values,
            "required": "Yes" if is_required else "No",
            "description": description
        })
        
        # Recursively handle nested objects
        if field_type == "object" and "properties" in field_def:
            nested_rows = extract_field_info(field_def["properties"], full_path)
            rows.extend(nested_rows)
        
        # Handle array items
        if field_type == "array" and "items" in field_def:
            items_def = field_def["items"]
            if "properties" in items_def:
                nested_rows = extract_field_info(items_def["properties"], f"{full_path}[]")
                rows.extend(nested_rows)
    
    return rows


def create_schema_dataframe(schema: dict) -> pd.DataFrame:
    """Create a DataFrame from the schema."""
    rows = []
    
    # Main properties
    properties = schema.get("properties", {})
    
    # Permit Details section
    if "permitDetails" in properties:
        rows.append({
            "Field Name": "PERMIT DETAILS",
            "Variable Path": "",
            "Type": "",
            "Required": "",
            "Enum Values": "",
            "Description": "",
            "_section": "section_header"
        })
        
        permit_props = properties["permitDetails"].get("properties", {})
        permit_required = properties["permitDetails"].get("required", [])
        
        for field_name, field_def in permit_props.items():
            # Get type
            field_type = ""
            if "type" in field_def:
                field_type = field_def["type"]
            elif "anyOf" in field_def:
                types = [t.get("type", "") for t in field_def["anyOf"] if "type" in t]
                field_type = " | ".join(filter(None, types))
            
            # Get enum values
            enum_values = ""
            if "enum" in field_def:
                enum_values = ", ".join([f'"{v}"' for v in field_def["enum"]])
            
            rows.append({
                "Field Name": field_name,
                "Variable Path": f"permitDetails.{field_name}",
                "Type": field_type,
                "Required": "Yes" if field_name in permit_required else "No",
                "Enum Values": enum_values,
                "Description": field_def.get("description", ""),
                "_section": "data"
            })
    
    # Generator Sets section
    if "generatorSets" in properties:
        rows.append({
            "Field Name": "GENERATOR SETS",
            "Variable Path": "",
            "Type": "",
            "Required": "",
            "Enum Values": "",
            "Description": properties["generatorSets"].get("description", ""),
            "_section": "section_header"
        })
        
        gen_items = properties["generatorSets"].get("items", {})
        gen_props = gen_items.get("properties", {})
        gen_required = gen_items.get("required", [])
        
        for field_name, field_def in gen_props.items():
            # Get type
            field_type = ""
            if "type" in field_def:
                field_type = field_def["type"]
            elif "anyOf" in field_def:
                types = [t.get("type", "") for t in field_def["anyOf"] if "type" in t]
                field_type = " | ".join(filter(None, types))
            
            # Get enum values
            enum_values = ""
            if "enum" in field_def:
                enum_values = ", ".join([f'"{v}"' for v in field_def["enum"]])
            
            rows.append({
                "Field Name": field_name,
                "Variable Path": f"generatorSets[].{field_name}",
                "Type": field_type,
                "Required": "Yes" if field_name in gen_required else "No",
                "Enum Values": enum_values,
                "Description": field_def.get("description", ""),
                "_section": "data"
            })
    
    return pd.DataFrame(rows)


def apply_styling(worksheet, df):
    """Apply professional styling to the worksheet."""
    # Define styles
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    
    section_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    section_font = Font(bold=True, color="FFFFFF", size=11)
    
    field_name_fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
    field_name_font = Font(bold=True, size=10)
    
    alt_row_light = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    alt_row_dark = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    
    # Borders
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )
    
    thick_border = Border(
        left=Side(style="medium"), right=Side(style="medium"),
        top=Side(style="medium"), bottom=Side(style="medium")
    )
    
    section_border = Border(
        left=Side(style="medium"), right=Side(style="medium"),
        top=Side(style="thick"), bottom=Side(style="medium")
    )
    
    center_align = Alignment(horizontal="center", vertical="top", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
    
    # Track data row counter for alternating colors
    data_row_counter = 0
    
    # Apply column headers (row 1)
    for col_num in range(1, 7):
        cell = worksheet.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thick_border
    
    # Apply row styles
    for row_num in range(2, len(df) + 2):
        section_type = df.iloc[row_num - 2].get("_section", "data")
        
        # Reset counter at section headers
        if section_type == "section_header":
            data_row_counter = 0
        elif section_type == "data":
            data_row_counter += 1
        
        use_alt_color = (data_row_counter % 2 == 0)
        
        for col_num in range(1, 7):
            cell = worksheet.cell(row=row_num, column=col_num)
            cell.alignment = left_align
            cell.border = thin_border
            
            if section_type == "section_header":
                # Section headers
                cell.fill = section_fill
                cell.font = section_font
                cell.border = section_border
                
            elif section_type == "data":
                # Data rows with alternating colors
                if col_num == 1:
                    # Field Name column - always gray
                    cell.fill = field_name_fill
                    cell.font = field_name_font
                else:
                    # Other columns - alternating white/light gray
                    cell.fill = alt_row_dark if use_alt_color else alt_row_light
    
    # Set column widths
    worksheet.column_dimensions["A"].width = 35  # Field Name
    worksheet.column_dimensions["B"].width = 40  # Variable Path
    worksheet.column_dimensions["C"].width = 15  # Type
    worksheet.column_dimensions["D"].width = 10  # Required
    worksheet.column_dimensions["E"].width = 30  # Enum Values
    worksheet.column_dimensions["F"].width = 80  # Description
    
    # Freeze the header row
    worksheet.freeze_panes = "A2"


def main():
    print("Creating schema documentation Excel file...")
    
    # Load schema
    print(f"\nLoading schema from: {SCHEMA_FILE}")
    schema = load_schema()
    
    # Create DataFrame
    print("Extracting field information...")
    df = create_schema_dataframe(schema)
    
    # Remove _section column for display
    df_display = df.drop(columns=["_section"])
    
    print(f"Found {len(df_display)} fields/sections")
    
    # Create output file
    output_path = WORKSPACE_ROOT / "schemas" / "air_quality_permits_schema_documentation.xlsx"
    
    print(f"\nWriting to: {output_path}")
    
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df_display.to_excel(writer, sheet_name="Schema Fields", index=False)
        
        # Apply styling
        worksheet = writer.sheets["Schema Fields"]
        apply_styling(worksheet, df)
    
    print(f"\n✅ Created: {output_path}")
    print("\nStructure:")
    print("  - Field Name: The variable name in the schema")
    print("  - Variable Path: The JSON path to access the field")
    print("  - Type: Data type (string, number, boolean, etc.)")
    print("  - Required: Whether the field is required")
    print("  - Enum Values: Allowed values if field is an enum")
    print("  - Description: Detailed description from the schema")
    print("\nColor coding:")
    print("  - Headers: Dark blue")
    print("  - Section headers: Blue")
    print("  - Field names: Gray")
    print("  - Data rows: Alternating white/light gray")


if __name__ == "__main__":
    main()
