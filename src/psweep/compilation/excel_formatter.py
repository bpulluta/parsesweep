#!/usr/bin/env python3
"""
Excel export with professional formatting.

Creates clean, readable Excel spreadsheets with proper styling.
"""

import pandas as pd
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class ExcelFormatter:
    """
    Export DataFrame to professionally formatted Excel file.

    Features:

    - Dark blue header with white bold text
    - Alternating row colors (white/light gray) for readability
    - Text wrapping enabled for long content
    - Centered alignment for short fields, left-aligned for details
    - Auto-adjusted column widths with smart limits
    - Frozen header row for scrolling
    - Clean borders throughout
    """

    @staticmethod
    def _sanitize_for_excel(value):
        """Remove Excel-illegal characters from cell values."""
        if not isinstance(value, str):
            return value
        # Remove control characters (U+0000 to U+001F, except tab/newline/cr)
        # Excel doesn't allow most control characters
        return "".join(
            ch if ord(ch) >= 32 or ch in "\t\n\r" else ""
            for ch in value
        )
    
    def save(
        self,
        df: pd.DataFrame,
        output_path: Path,
        freeze_columns: int = 0,
        auto_width: bool = True,
    ):
        """
        Save DataFrame to professionally formatted Excel file.

        Parameters
        ----------
        df : pd.DataFrame
            DataFrame to export
        output_path : Path
            Path for output Excel file
        freeze_columns : int
            Number of leading columns to freeze
        auto_width : bool
            Whether to auto-size columns

        Examples
        --------
        >>> formatter = ExcelFormatter()
        >>> formatter.save(df, Path("output.xlsx"))
        """
        output_path.parent.mkdir(parents=True, exist_ok=True)

        # Sanitize all string values to remove Excel-illegal characters
        df_clean = df.copy()
        for col in df_clean.columns:
            df_clean[col] = df_clean[col].apply(self._sanitize_for_excel)

        # Write data
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            df_clean.to_excel(writer, sheet_name="Data", index=False)

        # Load workbook for styling
        wb = load_workbook(output_path)
        ws = wb["Data"]

        # Apply all formatting
        self._apply_header_style(ws)
        center_aligned_cols = self._identify_center_aligned_columns(df, ws)
        self._apply_data_row_styles(ws, center_aligned_cols)
        if auto_width:
            self._apply_column_widths(ws, center_aligned_cols)
        self._freeze_header(ws, freeze_columns)

        # Save
        wb.save(output_path)

    def _apply_header_style(self, ws):
        """Apply dark blue header with white bold text."""
        header_fill = PatternFill(
            start_color="FF1F4E78",  # Dark blue
            end_color="FF1F4E78",
            fill_type="solid",
        )
        header_font = Font(bold=True, color="FFFFFFFF", size=11)
        thin_border = self._get_thin_border()

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.border = thin_border
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )

    def _identify_center_aligned_columns(self, df: pd.DataFrame, ws) -> list:
        """
        Identify columns that should be center-aligned.

        Short/categorical columns get centered, long text gets left-aligned.

        Parameters
        ----------
        df : pd.DataFrame
            Source DataFrame
        ws
            Worksheet object

        Returns
        -------
        list
            List of column letters that should be centered
        """
        center_aligned_cols = []

        for idx, col_name in enumerate(df.columns, start=1):
            col_letter = get_column_letter(idx)
            sample_values = df[col_name].dropna().astype(str).head(20)
            avg_length = (
                sample_values.str.len().mean() if len(sample_values) > 0 else 0
            )

            # Center if: numeric-like, dates, short categorical values
            if avg_length < 30 or col_name.lower() in [
                "state",
                "county",
                "city",
                "type",
                "value",
                "unit",
                "date",
                "number",
            ]:
                center_aligned_cols.append(col_letter)

        return center_aligned_cols

    def _apply_data_row_styles(self, ws, center_aligned_cols: list):
        """
        Apply alternating row colors and cell formatting.

        Parameters
        ----------
        ws
            Worksheet object
        center_aligned_cols : list
            List of column letters to center-align
        """
        # Define fills
        white_fill = PatternFill(
            start_color="FFFFFFFF", end_color="FFFFFFFF", fill_type="solid"
        )
        gray_fill = PatternFill(
            start_color="FFF2F2F2",  # Light gray
            end_color="FFF2F2F2",
            fill_type="solid",
        )

        data_font = Font(size=10)
        thin_border = self._get_thin_border()

        # Style data rows with alternating colors
        for row_idx, row in enumerate(
            ws.iter_rows(min_row=2, max_row=ws.max_row), start=2
        ):
            # Alternate row color
            row_fill = white_fill if row_idx % 2 == 0 else gray_fill

            for cell in row:
                cell.fill = row_fill
                cell.border = thin_border
                cell.font = data_font

                # Smart alignment based on column
                col_letter = get_column_letter(cell.column)
                if col_letter in center_aligned_cols:
                    cell.alignment = Alignment(
                        horizontal="center", vertical="top", wrap_text=True
                    )
                else:
                    cell.alignment = Alignment(
                        horizontal="left", vertical="top", wrap_text=True
                    )

    def _apply_column_widths(self, ws, center_aligned_cols: list):
        """
        Auto-adjust column widths with smart sizing.

        Parameters
        ----------
        ws
            Worksheet object
        center_aligned_cols : list
            List of column letters that are centered
        """
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)

            for cell in column:
                try:
                    cell_value = (
                        str(cell.value) if cell.value is not None else ""
                    )
                    if len(cell_value) > max_length:
                        max_length = len(cell_value)
                except Exception:  # noqa: BLE001 - column width is cosmetic; any cell error is non-fatal
                    pass

            # Set width with reasonable limits
            # Narrower for short columns, wider for details
            if column_letter in center_aligned_cols:
                adjusted_width = min(
                    max(max_length + 2, 12), 25
                )  # Centered cols: 12-25
            else:
                adjusted_width = min(
                    max(max_length + 2, 20), 60
                )  # Text cols: 20-60

            ws.column_dimensions[column_letter].width = adjusted_width

        # Set default row height to accommodate wrapped text
        for row in range(2, ws.max_row + 1):
            ws.row_dimensions[row].height = None  # Auto-height

    def _freeze_header(self, ws, freeze_columns: int = 0):
        """Freeze the header row and an optional number of leading columns."""
        freeze_columns = max(freeze_columns, 0)
        freeze_column_letter = get_column_letter(freeze_columns + 1)
        ws.freeze_panes = f"{freeze_column_letter}2"

    def _get_thin_border(self) -> Border:
        """Get standard thin border style."""
        return Border(
            left=Side(style="thin", color="D3D3D3"),
            right=Side(style="thin", color="D3D3D3"),
            top=Side(style="thin", color="D3D3D3"),
            bottom=Side(style="thin", color="D3D3D3"),
        )
