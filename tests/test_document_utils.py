from pathlib import Path

import openpyxl
import pytest

from psweep.extraction.document_utils import (
    _extract_from_docx,
    _extract_from_xlsx,
    _HTMLTextExtractor,
    extract_text_from_document,
    is_supported_document,
)


def test_extract_text_from_html_document(tmp_path: Path):
    html_path = tmp_path / "filing.html"
    html_path.write_text(
        """
<html>
  <head>
    <title>Apple Filing Summary</title>
    <style>.hidden { display:none; }</style>
  </head>
  <body>
    <h1>Apple Q1 2026</h1>
    <p>Net sales were $143.8 billion.</p>
    <script>console.log('ignore me')</script>
    <p>Net income was $42.1 billion.</p>
  </body>
</html>
""",
        encoding="utf-8",
    )

    extracted = extract_text_from_document(html_path)

    assert is_supported_document(html_path) is True
    assert "Apple Filing Summary" in extracted
    assert "Apple Q1 2026" in extracted
    assert "Net sales were $143.8 billion." in extracted
    assert "Net income was $42.1 billion." in extracted
    assert "ignore me" not in extracted


# ---------------------------------------------------------------------------
# Serializer output locks
#
# document_utils has two table serializers that are intentionally DIFFERENT
# from each other and from section_locator's bordered `| a | b |` format:
#   - `_extract_from_docx` (#3) : unbordered `a | b`, empty cells DROPPED
#   - `_HTMLTextExtractor` (#4) : tab-separated columns, empty cells kept
# These tests pin the exact byte output so the "kept separate" decision cannot
# regress into accidental drift.
# ---------------------------------------------------------------------------


class TestDocxExtractTableExactOutput:
    """Lock #3: unbordered `a | b`, empty cells dropped, all-blank rows dropped."""

    def test_exact_output(self, tmp_path: Path):
        import docx

        d = docx.Document()
        d.add_paragraph("Rate Schedule")
        tbl = d.add_table(rows=3, cols=2)
        tbl.rows[0].cells[0].text = "A"
        tbl.rows[0].cells[1].text = "B"
        tbl.rows[1].cells[0].text = "1"
        tbl.rows[1].cells[1].text = ""  # blank cell -> dropped entirely
        # row 2 left blank -> dropped
        path = tmp_path / "t.docx"
        d.save(path)

        # `Rate Schedule` is a plain paragraph; table rows follow, unbordered,
        # with the blank cell in row 1 omitted (`1`, not `1 | `).
        assert _extract_from_docx(path) == "Rate Schedule\nA | B\n1"


class TestHtmlExtractorTableExactOutput:
    """Lock #4: tab-separated columns, empty cells preserved."""

    def test_exact_output(self):
        html = (
            "<html><body><table>"
            "<tr><td>AG</td><td>500</td><td>feet</td></tr>"
            "<tr><td>R1</td><td></td><td>x</td></tr>"
            "</table></body></html>"
        )
        ex = _HTMLTextExtractor()
        ex.feed(html)
        assert ex.get_text() == "AG\t500\tfeet\n\nR1\t\tx"


def test_xlsx_workbook_closed_on_success_and_failure(tmp_path, monkeypatch):
    """Workbook handles close after serialization and row failures."""

    class Workbook:
        sheetnames = ["Sheet1"]
        closed = False
        fail = False

        def __getitem__(self, key):
            return self

        def iter_rows(self, **kwargs):
            if self.fail:
                raise OSError("bad row")
            return iter([(1, None, "value")])

        def close(self):
            self.closed = True

    wb = Workbook()
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *_a, **_k: wb)
    assert (
        _extract_from_xlsx(tmp_path / "a.xlsx")
        == "\n### Sheet: Sheet1 ###\n\n1 |  | value"
    )
    assert wb.closed
    wb.closed, wb.fail = False, True
    with pytest.raises(RuntimeError, match="bad row"):
        _extract_from_xlsx(tmp_path / "a.xlsx")
    assert wb.closed


def test_document_pdf_return_meta_passes_actual_scope(tmp_path, monkeypatch):
    """Document parsing exposes the parser's bounded actual PDF scope."""
    path = tmp_path / "a.pdf"
    path.write_bytes(b"fake")
    monkeypatch.setattr(
        "psweep.extraction.pdf_utils.extract_text_from_pdf",
        lambda *_a, **_k: (
            "selected",
            {"used_ocr": False, "page_range": (2, 3), "page_count": 3},
        ),
    )
    assert extract_text_from_document(path, (2, 9), return_meta=True) == (
        "selected",
        {"used_ocr": False, "page_range": (2, 3), "page_count": 3},
    )
